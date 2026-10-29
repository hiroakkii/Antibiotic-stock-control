"""内用薬・外用薬リストと薬効分類マスタから、アプリ用の医薬品マスタ(Excel)を作る。

使い方:
  python3 tools/build_master.py <内用薬.xlsx> <外用薬.xlsx> <薬効分類マスタを含む既存ブック.xlsx> <出力.xlsx>

出力はそのままでは数式の計算結果が入っていないため、LibreOffice で開いて保存し直す
(soffice --headless --convert-to xlsx) と、アプリや Excel 以外のツールからも値を読める。
"""
import sys
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

naiyo, gaiyo, classbook, out = sys.argv[1:5]

# 対象区分の初期値（分類コード → 区分）。Excel 上の「薬効分類マスタ」シートで自由に変更できる。
ANTIBIOTIC = {611, 612, 613, 614, 615, 619, 621, 624}
ANALGESIC = {114}
DEFAULT_YEARS = 3  # 仮値。製品の有効期限に合わせて変更すること

# ---------- 取り込み ----------
def read_list(path):
    ws = openpyxl.load_workbook(path, read_only=True)['ＨＰ用']
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[1]:
            continue
        brand = '先発品' if r[10] else ('後発品' if r[9] else '')
        rows.append([r[0], r[1], r[2], r[3], r[7], r[8], r[12], brand])
    return rows

drugs = read_list(naiyo) + read_list(gaiyo)
classes = [r for r in openpyxl.load_workbook(classbook, read_only=True)['薬効分類マスタ'].iter_rows(min_row=2, values_only=True) if r[0]]
N = len(drugs)
LAST = N + 1
NC = len(classes)
CLAST = NC + 1

wb = openpyxl.Workbook()
s_search = wb.active
s_search.title = '検索'
s_class = wb.create_sheet('薬効分類マスタ')
s_drug = wb.create_sheet('医薬品マスタ')

HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
CALC_FILL = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='999999')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOLD = Font(bold=True)

def head(ws, row, values, fill=HEAD_FILL):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = BOLD
        c.fill = fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX

# ---------- 薬効分類マスタ ----------
head(s_class, 1, ['薬効分類コード', '薬効分類名', '内用薬の件数', '外用薬の件数', '対象区分', '有効期限(年)', '検索一致(計算)', '検索順位(計算)'])
for c in (7, 8):
    s_class.cell(row=1, column=c).fill = CALC_FILL
for i, (code, name) in enumerate(classes, 2):
    s_class.cell(row=i, column=1, value=int(code))
    s_class.cell(row=i, column=2, value=name)
    s_class.cell(row=i, column=3, value=f'=COUNTIFS(医薬品マスタ!$C$2:$C${LAST},$A{i},医薬品マスタ!$A$2:$A${LAST},"内用薬")')
    s_class.cell(row=i, column=4, value=f'=COUNTIFS(医薬品マスタ!$C$2:$C${LAST},$A{i},医薬品マスタ!$A$2:$A${LAST},"外用薬")')
    kind = '抗菌薬' if int(code) in ANTIBIOTIC else '解熱鎮痛薬' if int(code) in ANALGESIC else None
    e = s_class.cell(row=i, column=5, value=kind)
    f = s_class.cell(row=i, column=6, value=DEFAULT_YEARS if kind else None)
    for c in (e, f):
        c.fill = INPUT_FILL
    s_class.cell(row=i, column=7, value=(
        f'=IF(検索!$B$25="",0,IF(OR(ISNUMBER(SEARCH(検索!$B$25,$A{i}&"")),ISNUMBER(SEARCH(検索!$B$25,$B{i}))),1,0))'
    ))
    s_class.cell(row=i, column=8, value=f'=N(H{i-1})+G{i}')
    for c in (7, 8):
        s_class.cell(row=i, column=c).fill = CALC_FILL
dv = DataValidation(type='list', formula1='"抗菌薬,解熱鎮痛薬"', allow_blank=True)
s_class.add_data_validation(dv)
dv.add(f'E2:E{CLAST}')
for col, w in zip('ABCDEFGH', [14, 56, 12, 12, 14, 12, 14, 14]):
    s_class.column_dimensions[col].width = w
s_class.freeze_panes = 'C2'
s_class.auto_filter.ref = f'A1:F{CLAST}'
s_class['J1'] = '【使い方】黄色いセルが設定欄です。「対象区分」に 抗菌薬 / 解熱鎮痛薬 を選ぶと、その薬効分類の医薬品が在庫管理アプリの薬剤マスタとして取り込まれます（空欄の分類は取り込まれません）。'
s_class['J2'] = '「有効期限(年)」は製造からの有効期限で、薬効分類ごとの標準値です。個別の医薬品だけ変えたい場合は「医薬品マスタ」シートの「個別設定(年)」に入力します。'
s_class['J3'] = '※ 初期設定の分類と有効期限(3年)は仮の値です。必ず確認してから使用してください。'
s_class['J3'].font = Font(bold=True, color='C00000')
s_class.column_dimensions['J'].width = 100

# ---------- 医薬品マスタ ----------
head(s_drug, 1, ['区分', '薬価基準収載医薬品コード', '薬効分類コード', '薬効分類名', '成分名', '規格', '品名', 'メーカー名', '薬価',
                 '先発/後発', '対象区分', '個別設定(年)', '有効期限(年)', '品名(検索用・半角)', '成分名(検索用・半角)', '一覧用(計算)'])
for c in (14, 15, 16):
    s_drug.cell(row=1, column=c).fill = CALC_FILL
s_drug.cell(row=1, column=12).fill = INPUT_FILL
for i, (kind, code, comp, spec, name, maker, price, brand) in enumerate(drugs, 2):
    s_drug.cell(row=i, column=1, value=kind)
    s_drug.cell(row=i, column=2, value=code)
    s_drug.cell(row=i, column=3, value=int(code[:3]))
    s_drug.cell(row=i, column=4, value=f'=IFERROR(VLOOKUP($C{i},薬効分類マスタ!$A$2:$B${CLAST},2,FALSE),"未登録")')
    s_drug.cell(row=i, column=5, value=comp)
    s_drug.cell(row=i, column=6, value=spec)
    s_drug.cell(row=i, column=7, value=name)
    s_drug.cell(row=i, column=8, value=maker)
    s_drug.cell(row=i, column=9, value=price)
    s_drug.cell(row=i, column=10, value=brand)
    s_drug.cell(row=i, column=11, value=f'=IFERROR(IF(VLOOKUP($C{i},薬効分類マスタ!$A$2:$E${CLAST},5,FALSE)=0,"",VLOOKUP($C{i},薬効分類マスタ!$A$2:$E${CLAST},5,FALSE)),"")')
    s_drug.cell(row=i, column=12).fill = INPUT_FILL
    s_drug.cell(row=i, column=13, value=f'=IF($K{i}="","",IF($L{i}<>"",$L{i},VLOOKUP($C{i},薬効分類マスタ!$A$2:$F${CLAST},6,FALSE)))')
    s_drug.cell(row=i, column=14, value=f'=ASC(G{i})')
    s_drug.cell(row=i, column=15, value=f'=ASC(E{i})')
    s_drug.cell(row=i, column=16, value=f'=N(P{i-1})+IF($C{i}=検索!$B$49,1,0)')
for col, w in zip('ABCDEFGHIJKLMNOP', [8, 16, 11, 34, 26, 18, 40, 20, 9, 9, 11, 11, 11, 24, 24, 10]):
    s_drug.column_dimensions[col].width = w
s_drug.freeze_panes = 'E2'
s_drug.auto_filter.ref = f'A1:M{LAST}'

# ---------- 検索 ----------
ws = s_search
ws.sheet_view.showGridLines = False
for col, w in zip('ABCDEFGH', [26, 50, 16, 16, 14, 4, 4, 4]):
    ws.column_dimensions[col].width = w

def title(r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(bold=True, size=14)

def note(r, text, rows=2):
    ws.merge_cells(start_row=r, start_column=1, end_row=r + rows - 1, end_column=5)
    c = ws.cell(row=r, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical='top')

def label(r, text):
    c = ws.cell(row=r, column=1, value=text)
    c.font = BOLD
    c.fill = HEAD_FILL
    c.border = BOX

def inbox(addr):
    ws[addr].fill = INPUT_FILL
    ws[addr].border = BOX
    ws[addr].font = Font(size=12, bold=True)

# 1) 医薬品名 → 薬効分類
title(1, '① 医薬品名 → 薬効分類名 検索')
note(2, '黄色いセルに医薬品名（品名または成分名。部分入力も可）を入力すると、薬価基準収載医薬品コードの左3桁から薬効分類名を表示します。', 2)
label(5, '医薬品名を入力'); inbox('B5')
ws.merge_cells('B5:E5')
label(7, '薬効分類名')
ws.merge_cells('B7:E7')
ws['B7'] = ('=IF($B$5="","",IF($H$5="","該当する医薬品が見つかりませんでした。名称の一部を短くするか、別の表記でお試しください。",'
            'IFERROR(INDEX(薬効分類マスタ!$B$2:$B$%d,MATCH($B$14,薬効分類マスタ!$A$2:$A$%d,0)),"分類コードが薬効分類マスタに未登録です")))' % (CLAST, CLAST))
ws['B7'].font = Font(bold=True, size=12, color='1F4E79')
rows = [(10, '該当した品名', 'G'), (11, '規格', 'F'), (12, '成分名', 'E'), (13, '薬価基準収載医薬品コード', 'B'), (15, '区分（内用薬/外用薬）', 'A'), (16, '薬価', 'I'), (17, '対象区分（アプリ用）', 'K')]
for r, lab, col in rows:
    label(r, lab)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.cell(row=r, column=2, value=f'=IF($H$5="","",INDEX(医薬品マスタ!${col}$2:${col}${LAST},$H$5))')
label(14, '薬効分類コード（左3桁）')
ws.merge_cells('B14:E14')
ws['B14'] = '=IF($B$13="","",VALUE(LEFT($B$13,3)))'
ws['B14'].alignment = Alignment(horizontal='left')
ws['H2'] = f'=IFERROR(MATCH("*"&ASC($B$5)&"*",医薬品マスタ!$N$2:$N${LAST},0),"")'
ws['H3'] = f'=IFERROR(MATCH("*"&ASC($B$5)&"*",医薬品マスタ!$O$2:$O${LAST},0),"")'
ws['H5'] = '=IF($B$5="","",IF(H2<>"",H2,IF(H3<>"",H3,"")))'
ws['G1'] = '内部計算用（編集・削除しないでください）'
ws['G1'].font = Font(color='999999', size=8)
note(19, '※ 同じ名称で複数の医薬品が存在する場合は、最初に見つかった1件を表示します。', 1)

# 2) 薬効分類から探す
title(23, '② 薬効分類を探す（キーワード or コード）')
note(24, '黄色いセルに薬効分類名の一部（例: 抗菌、グラム、解熱）または分類コード（例: 61、114）を入力すると、該当する薬効分類を最大15件まで表示します。', 1)
label(25, 'キーワード/コード'); inbox('B25')
ws.merge_cells('B25:E25')
head(ws, 27, ['薬効分類コード', '薬効分類名', '内用薬の件数', '外用薬の件数', '対象区分'])
for k in range(1, 16):
    r = 27 + k
    m = f'MATCH({k},薬効分類マスタ!$H$2:$H${CLAST},0)'
    for col, src in zip(range(1, 6), 'ABCDE'):
        c = ws.cell(row=r, column=col, value=f'=IFERROR(INDEX(薬効分類マスタ!${src}$2:${src}${CLAST},{m})&"","")' if col == 5 else f'=IFERROR(INDEX(薬効分類マスタ!${src}$2:${src}${CLAST},{m}),"")')
        c.border = BOX
        if col != 2:
            c.alignment = Alignment(horizontal='center')
ws['A43'] = '=IF($B$25="","",IF(MAX(薬効分類マスタ!$H$2:$H$%d)=0,"該当する薬効分類がありません",IF(MAX(薬効分類マスタ!$H$2:$H$%d)>15,"※ "&MAX(薬効分類マスタ!$H$2:$H$%d)&"件該当しました（先頭15件を表示）","")))' % (CLAST, CLAST, CLAST)
ws['A43'].font = Font(color='C00000')

# 3) 薬効分類コードから医薬品を一覧
title(46, '③ 薬効分類コードから医薬品を一覧')
note(47, '黄色いセルに薬効分類コード3桁（例: 114）を入力すると、該当する医薬品（内用薬・外用薬）を最大30件まで表示します。全件は「医薬品マスタ」シートで「薬効分類コード」を絞り込んでください。', 2)
label(49, '薬効分類コード（3桁）'); inbox('B49')
ws.merge_cells('B49:E49')
ws['B49'].alignment = Alignment(horizontal='left')
head(ws, 51, ['区分／成分名', '品名', '規格', '薬価', '対象区分'])
for k in range(1, 31):
    r = 51 + k
    m = f'MATCH({k},医薬品マスタ!$P$2:$P${LAST},0)'
    def idx(col):
        return f'INDEX(医薬品マスタ!${col}$2:${col}${LAST},{m})'
    ws.cell(row=r, column=1, value=f'=IFERROR({idx("A")}&"／"&{idx("E")},"")')
    ws.cell(row=r, column=2, value=f'=IFERROR({idx("G")},"")')
    ws.cell(row=r, column=3, value=f'=IFERROR({idx("F")},"")')
    ws.cell(row=r, column=4, value=f'=IFERROR({idx("I")},"")')
    ws.cell(row=r, column=5, value=f'=IFERROR({idx("K")},"")')
    for c in range(1, 6):
        ws.cell(row=r, column=c).border = BOX
ws['A83'] = f'=IF(ISNUMBER($B$49),IF(MAX(医薬品マスタ!$P$2:$P${LAST})=0,"該当する医薬品がありません",IF(MAX(医薬品マスタ!$P$2:$P${LAST})>30,"※ "&MAX(医薬品マスタ!$P$2:$P${LAST})&"件該当しました（先頭30件を表示）","")),"")'
ws['A83'].font = Font(color='C00000')
dv2 = DataValidation(type='whole', operator='between', formula1='100', formula2='999', allow_blank=True,
                     errorTitle='入力エラー', error='薬効分類コードは3桁の数字で入力してください。')
ws.add_data_validation(dv2)
dv2.add('B49')

wb.save(out)
print('drugs', N, 'classes', NC)
